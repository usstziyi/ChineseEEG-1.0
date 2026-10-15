"""从 segmented_novel 下的整本小说文件中提取逐行文本，导出为 CSV。

数据源约定
----------
正文文本只取自 ``segmented_Chinense_novel.xlsx``（单列 Chinese_text），
它是「前言 + 全部正文」的完整版本；``_display.xlsx`` 是它的逐帧重复副本，
不参与提取。

``run_N.xlsx`` 本身不提供文本，只用来确定 run 边界：已验证
「主文件正文 == 按数字序拼接 run_1..run_N」，且每个 run 起点都落在章节号行上
（GarnettDream 1 run = 1 章，LittlePrince 1 run = 4 章），
所以 run 划分只能由这些文件的长度推出。

处理规则
--------
1. 前言：首个章节号之前的行全部丢弃。
2. 章节号：形如 ``19`` 的纯数字行只用于标记章节起点，不写入正文。
3. ``run``：EEG 采集轮次，从 1 起编号。
4. ``sentence`` 为原始折行片段，不做拼接或断句处理。
5. ``n_chars``：非标点字符数，标点按 Unicode 类别 ``P*`` 判定。
"""

from __future__ import annotations

import csv
import re
import unicodedata
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SEGMENTED_ROOT = PROJECT_ROOT / "datasets/ChineseEEG/derivatives/novels/segmented_novel"
OUTPUT_DIR = PROJECT_ROOT / "datasets/preprocessed/novels"

NOVEL_FILE = "segmented_Chinense_novel.xlsx"
TEXT_COLUMN = 0  # 单列 Chinese_text
FIELDS = ["index", "chapter", "run", "sentence", "n_chars"]

CHAPTER_RE = re.compile(r"^\d+$")
RUN_FILE_RE = re.compile(rf"{Path(NOVEL_FILE).stem}_run_(\d+)\.xlsx$")


def count_chars(text: str) -> int:
    """统计非标点字符数，Unicode 类别以 P 开头的字符视为标点。"""
    return sum(1 for char in text if not unicodedata.category(char).startswith("P"))


def read_lines(xlsx_path: Path) -> list[str]:
    """读取单列 xlsx，返回去掉表头后每一行的文本。"""
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    lines = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True)):
        if row_number == 0:  # 表头 Chinese_text
            continue
        cell = row[TEXT_COLUMN]
        lines.append("" if cell is None else str(cell).strip())
    workbook.close()
    return lines


def read_run_lengths(novel_dir: Path) -> list[int]:
    """按 run 序号升序读取各 run 文件的行数。"""
    run_files = []
    for path in novel_dir.iterdir():
        matched = RUN_FILE_RE.match(path.name)
        if matched:
            run_files.append((int(matched.group(1)), path))
    run_files.sort()
    return [len(read_lines(path)) for _, path in run_files]


def extract(novel_dir: Path) -> list[dict]:
    """整理出逐行记录，返回 index / chapter / run / sentence / n_chars 列表。"""
    lines = read_lines(novel_dir / NOVEL_FILE)
    run_lengths = read_run_lengths(novel_dir)

    # 前言：首个章节号之前的内容整段跳过
    body_start = next(i for i, text in enumerate(lines) if CHAPTER_RE.match(text))

    # 按各 run 的行数把正文逐段划归到 run，run 序号从 1 起
    run_of_line = [0] * len(lines)
    cursor = body_start
    for run, length in enumerate(run_lengths, start=1):
        for i in range(cursor, cursor + length):
            run_of_line[i] = run
        cursor += length
    if cursor != len(lines):
        raise ValueError(f"{novel_dir.name}: run 文件合计 {cursor - body_start} 行，正文为 {len(lines) - body_start} 行")

    records: list[dict] = []
    chapter = 0
    for i in range(body_start, len(lines)):
        text = lines[i]
        if CHAPTER_RE.match(text):
            chapter = int(text)  # 章节号仅用于划分章节
            continue
        if not text:
            continue
        records.append(
            {
                "index": len(records),
                "chapter": chapter,
                "run": run_of_line[i],
                "sentence": text,
                "n_chars": count_chars(text),
            }
        )
    return records


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for novel_dir in sorted(path for path in SEGMENTED_ROOT.iterdir() if path.is_dir()):
        records = extract(novel_dir)
        csv_path = OUTPUT_DIR / f"{novel_dir.name}.csv"
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(records)
        print(
            f"{novel_dir.name}: {len(records)} 行, "
            f"章节 1-{max(r['chapter'] for r in records)}, "
            f"run 1-{max(r['run'] for r in records)} -> {csv_path}"
        )


if __name__ == "__main__":
    main()
